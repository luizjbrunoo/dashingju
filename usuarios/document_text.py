from pathlib import Path
from typing import Any

import pypdfium2 as pdfium
from docx import Document as DocxDocument


def extract_document_text(file_obj: Any, file_name: str = "") -> str:
    ext = Path(file_name or "").suffix.lower()

    try:
        if ext == ".txt":
            raw = file_obj.read()
            return raw.decode("utf-8", errors="ignore").strip()

        if ext == ".docx":
            doc = DocxDocument(file_obj)
            lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            return "\n".join(lines).strip()

        if ext == ".pdf":
            pdf = pdfium.PdfDocument(file_obj)
            pages_text = []
            for page in pdf:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                    if text and text.strip():
                        pages_text.append(text.strip())
                finally:
                    textpage.close()
                    page.close()
            pdf.close()
            return "\n\n".join(pages_text).strip()
    except Exception:
        return ""
    finally:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)

    return ""
